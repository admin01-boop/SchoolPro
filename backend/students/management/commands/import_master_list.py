"""Import the school's master/trial student-list spreadsheet into the database.

Usage:
    python manage.py import_master_list "path\\to\\Master List 2026-2027.xlsx" --academic-year "2026-2027"
"""
from datetime import date

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from academics.features import is_diploma_programme
from academics.models import (
    AcademicConfiguration,
    AcademicYear,
    ClassRoom,
    Curriculum,
    GradeLevel,
    LeadSource,
    Programme,
)
from students.models import EmergencyContact, Enrollment, Guardian, Student, StudentGuardian

MASTER_SHEET_NAME = 'MASTER LIST 2026-2027'
TRIAL_SHEET_NAME = 'TRIAL STUDENT LIST'
HEADER_ROW = 5
FIRST_DATA_ROW = 6

# Column letter -> field name, matching the workbook's own header row.
MASTER_COLUMNS = {
    'B': 'start_date', 'C': 'enrollment_status', 'D': 'frn', 'E': 'student_id',
    'F': 'full_name', 'G': 'khmer_name', 'H': 'sex', 'I': 'date_of_birth',
    'K': 'nationality_1', 'L': 'nationality_2', 'M': 'curriculum', 'N': 'programme',
    'O': 'grade_level', 'P': 'english_class', 'Q': 'khmer_class', 'R': 'learning_support',
    'S': 'father_name', 'T': 'father_phone_1', 'U': 'father_phone_2',
    'V': 'mother_name', 'W': 'mother_phone_1', 'X': 'mother_phone_2',
    'Y': 'legal_custody', 'Z': 'emergency_name', 'AA': 'emergency_relationship',
    'AB': 'emergency_phone', 'AC': 'remark',
}

TRIAL_COLUMNS = {
    'B': 'start_date', 'C': 'end_date', 'D': 'enrollment_status', 'E': 'frn', 'F': 'student_id',
    'G': 'full_name', 'H': 'khmer_name', 'I': 'sex', 'J': 'date_of_birth',
    'L': 'nationality_1', 'M': 'nationality_2', 'N': 'curriculum', 'O': 'programme',
    'P': 'grade_level', 'Q': 'english_class', 'R': 'khmer_class',
    'S': 'father_name', 'T': 'father_phone_1', 'U': 'mother_name', 'V': 'mother_phone_1',
    'W': 'legal_custody', 'X': 'lead_source',
}

STATUS_MAP = {
    'existing': Enrollment.Status.EXISTING,
    'new': Enrollment.Status.NEW,
    'new enrollment': Enrollment.Status.NEW,
    'returned': Enrollment.Status.RETURNED,
    'returned student': Enrollment.Status.RETURNED,
    'withdrawn': Enrollment.Status.WITHDRAWN,
    'left': Enrollment.Status.WITHDRAWN,
}


def clean(value):
    if value is None:
        return ''
    text = str(value).strip()
    if text.upper() in ('N/A', 'NA', '-'):
        return ''
    return text


def to_date(value):
    if isinstance(value, date):
        return value
    return None


class Command(BaseCommand):
    help = 'Import students from the school master/trial list Excel workbook.'

    def add_arguments(self, parser):
        parser.add_argument('file_path', type=str, help='Path to the .xlsx workbook.')
        parser.add_argument('--academic-year', type=str, default='2026-2027')
        parser.add_argument('--dry-run', action='store_true', help='Parse and validate without saving.')

    def handle(self, *args, **options):
        file_path = options['file_path']
        dry_run = options['dry_run']
        try:
            workbook = load_workbook(filename=file_path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f'File not found: {file_path}') from exc

        academic_year, _ = AcademicYear.objects.get_or_create(name=options['academic_year'])
        self.classes_enabled = AcademicConfiguration.load().classes_enabled

        self.unmapped_statuses = set()
        stats = {'master': 0, 'trial': 0}

        with transaction.atomic():
            if MASTER_SHEET_NAME in workbook.sheetnames:
                stats['master'] = self.import_sheet(
                    workbook[MASTER_SHEET_NAME], MASTER_COLUMNS, academic_year, is_trial=False
                )
            if TRIAL_SHEET_NAME in workbook.sheetnames:
                stats['trial'] = self.import_sheet(
                    workbook[TRIAL_SHEET_NAME], TRIAL_COLUMNS, academic_year, is_trial=True
                )
            if dry_run:
                transaction.set_rollback(True)

        if self.unmapped_statuses:
            self.stdout.write(self.style.WARNING(
                f'Unrecognized enrollment status values (defaulted to EXISTING): {sorted(self.unmapped_statuses)}'
            ))
        mode = 'DRY RUN - nothing saved' if dry_run else 'Imported'
        self.stdout.write(self.style.SUCCESS(
            f"{mode}: {stats['master']} master-list rows, {stats['trial']} trial-list rows."
        ))

    def import_sheet(self, sheet, columns, academic_year, is_trial):
        # sheet.max_row is unreliable here: the workbook's print-area/filter formatting
        # extends far past the real data, so we stop after a long run of blank rows instead.
        count = 0
        blank_streak = 0
        row_idx = FIRST_DATA_ROW
        max_row_idx = FIRST_DATA_ROW + 5000
        while blank_streak < 200 and row_idx < max_row_idx:
            row = {
                field: sheet[f'{col}{row_idx}'].value
                for col, field in columns.items()
            }
            student_id = clean(row.get('student_id'))
            full_name = clean(row.get('full_name'))
            if not student_id or not full_name:
                blank_streak += 1
                row_idx += 1
                continue
            blank_streak = 0
            self.import_row(row, academic_year, is_trial)
            count += 1
            row_idx += 1
        return count


    def import_row(self, row, academic_year, is_trial):
        student = self.upsert_student(row)
        curriculum = self.get_lookup(Curriculum, row.get('curriculum'))
        programme = self.get_lookup(Programme, row.get('programme'))
        grade_level = self.get_lookup(GradeLevel, row.get('grade_level'))
        can_change_class = self.classes_enabled or is_diploma_programme(programme)
        class_room = (
            self.get_classroom(academic_year, grade_level, row.get('english_class'), row.get('khmer_class'))
            if can_change_class else None
        )
        lead_source = self.get_lookup(LeadSource, row.get('lead_source')) if is_trial else None

        defaults = {
            'enrollment_status': self.map_status(row.get('enrollment_status')),
            'start_date': to_date(row.get('start_date')) or date.today(),
            'end_date': to_date(row.get('end_date')),
            'curriculum': curriculum,
            'programme': programme,
            'grade_level': grade_level,
            'learning_support_programme': clean(row.get('learning_support')),
            'lead_source': lead_source,
        }
        if can_change_class:
            defaults['class_room'] = class_room
        Enrollment.objects.update_or_create(
            student=student, academic_year=academic_year, is_trial=is_trial, defaults=defaults,
        )

        legal_custody = clean(row.get('legal_custody')).lower()
        self.upsert_guardian(
            student, clean(row.get('father_name')), clean(row.get('father_phone_1')),
            clean(row.get('father_phone_2')), StudentGuardian.Relationship.FATHER,
            is_legal_custody='father' in legal_custody or 'both' in legal_custody,
        )
        self.upsert_guardian(
            student, clean(row.get('mother_name')), clean(row.get('mother_phone_1')),
            clean(row.get('mother_phone_2')), StudentGuardian.Relationship.MOTHER,
            is_legal_custody='mother' in legal_custody or 'both' in legal_custody,
        )

        emergency_name = clean(row.get('emergency_name'))
        if emergency_name:
            EmergencyContact.objects.update_or_create(
                student=student, name=emergency_name,
                defaults={
                    'relationship': clean(row.get('emergency_relationship')),
                    'phone': clean(row.get('emergency_phone')),
                },
            )

    def upsert_student(self, row):
        sex_raw = clean(row.get('sex')).lower()
        sex = Student.Sex.FEMALE if sex_raw.startswith('f') else Student.Sex.MALE
        defaults = {
            'full_name': clean(row.get('full_name')),
            'khmer_name': clean(row.get('khmer_name')),
            'sex': sex,
            'date_of_birth': to_date(row.get('date_of_birth')) or date(2000, 1, 1),
            'nationality_1': clean(row.get('nationality_1')),
            'nationality_2': clean(row.get('nationality_2')),
        }
        frn = clean(row.get('frn'))
        if frn:
            defaults['frn'] = frn
        if 'remark' in row:
            defaults['remark'] = clean(row.get('remark'))
        student, _ = Student.objects.update_or_create(student_id=clean(row['student_id']), defaults=defaults)
        return student

    def upsert_guardian(self, student, name, phone_1, phone_2, relationship, is_legal_custody):
        if not name:
            return
        guardian, _ = Guardian.objects.get_or_create(
            full_name=name, phone_1=phone_1, defaults={'phone_2': phone_2}
        )
        StudentGuardian.objects.update_or_create(
            student=student, guardian=guardian, relationship=relationship,
            defaults={'is_legal_custody': is_legal_custody},
        )

    def get_lookup(self, model, name):
        name = clean(name)
        if not name:
            return None
        obj, _ = model.objects.get_or_create(name=name)
        return obj

    def get_classroom(self, academic_year, grade_level, english_name, khmer_name):
        english_name = clean(english_name)
        if not english_name or not grade_level:
            return None
        classroom, _ = ClassRoom.objects.get_or_create(
            academic_year=academic_year, grade_level=grade_level, english_name=english_name,
            defaults={'khmer_name': clean(khmer_name)},
        )
        return classroom

    def map_status(self, raw_value):
        text = clean(raw_value).lower()
        status = STATUS_MAP.get(text)
        if status is None:
            if text:
                self.unmapped_statuses.add(clean(raw_value))
            return Enrollment.Status.EXISTING
        return status
